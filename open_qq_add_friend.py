"""Search unique numbers from XLSX or ODS in QQ's Add Friends/Groups page."""

from __future__ import annotations

import re
import subprocess
import sys
import time
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from pathlib import Path
from task_control import atomic_save
from add_progress import AttemptJournal, ensure_number_file_writable
from gui_config import DEFAULT_VERIFICATION_MESSAGE
from qq_verification import (
    NoVerificationField,
    cancel_unconfigured_form,
    prepare_verification_form,
    click_verified_send,
)

# QQ 客户端位置、窗口标题和页面入口名称；客户端路径变化时修改 QQ_EXE。
QQ_EXE = Path(r"D:\01App\Tencent\QQNT\QQ.exe")
WINDOW_TITLE = "QQ"
PAGE_TITLE = "全网搜索"
MENU_ITEM_TEXT = "加好友/群"
# 搜索、页面加载和重试的等待参数，单位为秒（重试次数除外）。
STARTUP_TIMEOUT_SECONDS = 180
PAGE_TIMEOUT_SECONDS = 20
SEARCH_TIMEOUT_SECONDS = 20
MAX_SEARCH_ATTEMPTS = 3
RESULT_STABLE_SECONDS = 1.0
USER_TAB_TEXT = "用户"
STATIC_PAGE_TEXTS = {
    "全网搜索",
    "搜索",
    "全部",
    "用户",
    "群聊",
    "小程序",
    "机器人",
    "推荐",
    "游戏",
    "二次元",
    "学习",
    "音乐",
    "影视",
    "运动",
}
WORKSPACE_DIR = Path(__file__).resolve().parents[1]
# 支持 XLSX 或 ODS 完整路径；保留旧脚本默认值，桌面程序默认选择采集的 XLSX。
NUMBER_FILE_PATH = WORKSPACE_DIR / "号码.ods"
# 以下文本集合用于识别页面状态及可点击操作按钮。
NO_RESULT_MARKERS = (
    "没有找到",
    "未找到",
    "无搜索结果",
    "没有搜索结果",
    "未搜索到",
    "暂无搜索结果",
    "暂无相关结果",
    "没有相关结果",
    "没有相关用户",
    "没有相关群聊",
    "没有结果",
    "暂无结果",
)
LOADING_MARKERS = ("搜索中", "正在搜索", "加载中", "正在加载")
RESULT_ACTIONS = {"添加", "加入", "申请加入"}
SEND_BUTTON_TEXT = "发送"
AVATAR_BUTTON_TEXT = "昵称"
AVATAR_AUTO_DETECT = True
MORE_MENU_TEXT = "更多"
TARGET_PROCESS_ID = 0
AUTO_SEND = True
VERIFICATION_MESSAGE = DEFAULT_VERIFICATION_MESSAGE
NUMBER_INTERVAL_SECONDS = 1.0
SEND_TIMEOUT_SECONDS = 5.0
AVATAR_TIMEOUT_SECONDS = 5.0
CLOSE_SECONDARY_WINDOWS = True
# 每个参与轮换的账号本轮处理数量；修改此整数可调整每账号批次数量。
NUMBERS_PER_ACCOUNT = 3
# 按 QQ 号排除账号，例如 {"123456789", "987654321"}；空集合表示不排除。
EXCLUDED_QQ_ACCOUNTS = {}
# 写入号码表 B 列的完成标记，下次启动时据此跳过对应号码。
PROCESSED_MARKER = "已处理"


# 延迟导入桌面自动化和 ODS 依赖，缺少依赖时给出安装命令。
def load_dependencies():
    try:
        import pyautogui
        import win32con
        import win32gui
        from pywinauto import Desktop
        from odf.opendocument import load as load_ods
        from odf.table import Table, TableCell, TableRow
        from odf.teletype import extractText
        from odf.text import P
    except ImportError as error:
        raise RuntimeError(
            "缺少桌面自动化依赖，请运行："
            "python -m pip install pywinauto pywin32 pyautogui odfpy"
        ) from error
    return (
        Desktop,
        pyautogui,
        win32con,
        win32gui,
        load_ods,
        Table,
        TableRow,
        TableCell,
        extractText,
        P,
    )


# 选择一个可用 QQ 主窗口；优先选当前前台窗口，其次选面积较大的窗口。
def find_qq_window(desktop, win32gui):
    windows = find_qq_windows(desktop, win32gui)
    if not windows:
        return None
    foreground = win32gui.GetForegroundWindow()
    return max(
        windows,
        key=lambda window: (
            window.handle == foreground,
            window.rectangle().width() * window.rectangle().height(),
        ),
    )


# 枚举各个 QQ 登录实例的主窗口，并用进程号去重。
# 最小化窗口的屏幕矩形可能为零，因此改用窗口还原尺寸判断有效性。
def find_qq_windows(desktop, win32gui):
    windows = []
    process_ids = set()
    for window in desktop.windows():
        try:
            if window.window_text().strip() != WINDOW_TITLE:
                continue
            process_id = window.process_id()
            if TARGET_PROCESS_ID and process_id != TARGET_PROCESS_ID:
                continue
            handle = window.handle
            iconic = bool(win32gui.IsIconic(handle))
            if iconic:
                rect = win32gui.GetWindowPlacement(handle)[4]
                width = rect[2] - rect[0]
                height = rect[3] - rect[1]
            else:
                rect = window.rectangle()
                width = rect.right - rect.left
                height = rect.bottom - rect.top
            if (
                width <= 0
                or height <= 0
                or process_id in process_ids
                or (not win32gui.IsWindowVisible(handle) and not iconic)
            ):
                continue
            process_ids.add(process_id)
            windows.append(window)
        except Exception:
            continue
    return windows


# 返回识别到的 QQ 主窗口数量。
def count_qq_windows(desktop, win32gui):
    return len(find_qq_windows(desktop, win32gui))


# 关闭参与本次轮换账号的可见附属窗口，保留它们的主窗口。
def close_qq_secondary_windows(desktop, qq_windows, win32con, win32gui):
    main_handles = {window.handle for window in qq_windows}
    process_ids = {window.process_id() for window in qq_windows}
    closed_titles = []

    for window in desktop.windows():
        try:
            handle = window.handle
            if (
                window.process_id() not in process_ids
                or handle in main_handles
                or not win32gui.IsWindowVisible(handle)
            ):
                continue
            title = window.window_text().strip()
            if title in search_page_titles():
                # 保留已打开的搜索页面，后续直接复用，避免关闭后重新走菜单。
                continue
            win32gui.PostMessage(handle, win32con.WM_CLOSE, 0, 0)
            closed_titles.append(title or "未命名窗口")
        except Exception:
            continue

    if closed_titles:
        time.sleep(0.5)
        print("已关闭 QQ 主窗口之外的页面：" + "、".join(closed_titles))
    else:
        print("没有需要关闭的 QQ 其他页面。")


def search_page_titles():
    """兼容新旧 QQ 搜索窗口标题，保留用户配置的自定义标题。"""
    return {PAGE_TITLE, '全网搜索', '综合搜索'}


# 在指定 QQ 进程中查找已打开的新旧搜索页面。
def find_page_window(desktop, process_id):
    for window in desktop.windows():
        try:
            if (
                window.process_id() == process_id
                and window.window_text().strip() in search_page_titles()
                and window.is_visible()
            ):
                return window
        except Exception:
            continue
    return None


# 等待 QQ 主窗口出现，直到达到给定超时时间。
def wait_for_qq_window(desktop, win32gui, timeout):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        window = find_qq_window(desktop, win32gui)
        if window is not None:
            return window
        time.sleep(1)
    return None


# 在指定 QQ 实例的弹出菜单中定位“加好友/群”。
def find_menu_item(desktop, process_id):
    for window in desktop.windows():
        try:
            if window.process_id() != process_id:
                continue
            for control in window.descendants():
                if (
                    control.element_info.control_type in {"Button", "MenuItem", "Text"}
                    and control.window_text().strip() == MENU_ITEM_TEXT
                    and control.is_visible()
                    and control.is_enabled()
                ):
                    return control
        except Exception:
            continue
    return None


# 定位会话搜索框右侧的添加入口，不使用侧栏设置或聊天工具栏中的同名按钮。
def find_more_menu_item(desktop, process_id, window=None):
    windows = [window] if window is not None else desktop.windows()
    candidates = []
    semantic_matches = []
    for main in windows:
        try:
            if (main.process_id() != process_id or main.window_text().strip() != WINDOW_TITLE
                    or not main.is_visible()):
                continue
            bounds = main.rectangle()
            controls = main.descendants()
            searches = []
            more_buttons = []
            for control in controls:
                if not control.is_visible() or not control.is_enabled():
                    continue
                name = control.window_text().strip()
                kind = control.element_info.control_type
                area = control.rectangle()
                if (kind == 'Edit' and name == '搜索'
                        and bounds.left <= area.left < bounds.left + (bounds.right - bounds.left) * 0.55
                        and bounds.top <= area.top < bounds.top + min(220, (bounds.bottom - bounds.top) * 0.3)):
                    searches.append(control)
                if kind not in {'Button', 'MenuItem'} or name != MORE_MENU_TEXT:
                    continue
                if not (bounds.left <= area.left < area.right <= bounds.right
                        and bounds.top <= area.top < area.bottom <= bounds.bottom):
                    continue
                more_buttons.append(control)
                class_tokens = set((control.element_info.class_name or '').split())
                if class_tokens & {'contact-adder-btn', 'top-bar__adder'}:
                    semantic_matches.append(control)
            for control in more_buttons:
                area = control.rectangle()
                for search in searches:
                    search_area = search.rectangle()
                    overlap = min(area.bottom, search_area.bottom) - max(area.top, search_area.top)
                    height = min(area.bottom - area.top, search_area.bottom - search_area.top)
                    if height > 0 and overlap >= height * 0.5 and -3 <= area.left - search_area.right <= height * 3:
                        candidates.append(control)
                        break
        except Exception:
            continue
    if semantic_matches:
        return semantic_matches[0] if len(semantic_matches) == 1 else None
    return candidates[0] if len(candidates) == 1 else None


# 恢复并激活指定窗口，使后续 UI 自动化操作作用于正确窗口。
def activate_window(window, win32con, win32gui):
    handle = window.handle
    if win32gui.IsIconic(handle):
        win32gui.ShowWindow(handle, win32con.SW_RESTORE)
    win32gui.BringWindowToTop(handle)
    win32gui.SetForegroundWindow(handle)
    time.sleep(0.8)


# 切换账号时最小化其他主窗口，并确认目标窗口确实处于前台。
def activate_qq_account_window(window, qq_windows, win32con, win32gui):
    for other_window in qq_windows:
        if other_window.handle != window.handle:
            win32gui.ShowWindow(other_window.handle, win32con.SW_MINIMIZE)

    activate_window(window, win32con, win32gui)
    if win32gui.GetForegroundWindow() != window.handle:
        window.set_focus()
        time.sleep(0.5)
    if win32gui.GetForegroundWindow() != window.handle:
        raise RuntimeError(
            f"无法将 QQ 账号窗口 {window.process_id()} 置于前台，已停止以避免操作错账号。"
        )


# 打开指定账号的加好友/群全网搜索页；若页面已存在则直接复用。
def open_add_friend_page(desktop, win32con, win32gui, window=None):
    if window is None:
        window = find_qq_window(desktop, win32gui)
    if window is None:
        if not QQ_EXE.is_file():
            raise RuntimeError(f"找不到 QQ 程序：{QQ_EXE}，请修改脚本中的 QQ_EXE 路径。")
        subprocess.Popen([str(QQ_EXE)])
        window = wait_for_qq_window(desktop, win32gui, STARTUP_TIMEOUT_SECONDS)
        if window is None:
            raise RuntimeError(
                f"等待 QQ 主窗口超时（{STARTUP_TIMEOUT_SECONDS} 秒）。"
                "请确认 QQ 已启动并完成登录后重试。"
            )

    activate_window(window, win32con, win32gui)
    process_id = window.process_id()
    page = find_page_window(desktop, process_id)
    if page is not None:
        page.set_focus()
        print(f'复用已打开的 QQ 搜索页面：{page.window_text().strip()}。')
        return window, page

    rect = window.rectangle()
    width = rect.right - rect.left
    height = rect.bottom - rect.top
    if width < 240 or height < 300:
        raise RuntimeError("QQ 主窗口尺寸异常，无法安全打开搜索页面。")

    menu_item = find_menu_item(desktop, process_id)
    if menu_item is None:
        more_menu_item = find_more_menu_item(desktop, process_id, window)
        if more_menu_item is None:
            raise RuntimeError('未找到会话搜索框右侧的添加入口。请手动打开“加好友/群”搜索页面后重试。')
        more_menu_item.click_input()
        print('已点击会话搜索框右侧的添加入口。')
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            menu_item = find_menu_item(desktop, process_id)
            if menu_item is not None:
                break
            time.sleep(0.2)

    if menu_item is None:
        raise RuntimeError(
            "没有检测到“加好友/群”菜单项；为避免误点其他内容，脚本已停止。"
        )
    menu_item.click_input()

    deadline = time.monotonic() + PAGE_TIMEOUT_SECONDS
    while time.monotonic() < deadline:
        page = find_page_window(desktop, process_id)
        if page is not None:
            page.set_focus()
            print(f"已打开 QQ“{MENU_ITEM_TEXT}”页面（{page.window_text().strip()}）。")
            return window, page
        time.sleep(0.5)

    raise RuntimeError(
        '点击菜单后未检测到“综合搜索”或“全网搜索”窗口。'
        "请确认 QQ 使用中文界面且当前版本的菜单布局兼容。"
    )


def find_own_avatar(window):
    """只从账号栏选择头像，避免把聊天或成员头像当作当前登录账号。"""
    candidates = []
    account_header = []
    for control in window.descendants():
        try:
            name = control.window_text().strip()
            if (control.element_info.control_type != 'Button'
                    or not control.is_visible() or not control.is_enabled()
                    or not (name == AVATAR_BUTTON_TEXT
                            or (AVATAR_AUTO_DETECT and name.endswith('的头像')))):
                continue
            candidates.append(control)
            if any(sibling.is_visible() and sibling.window_text().strip() == '在线状态'
                   for sibling in control.parent().children()):
                account_header.append(control)
        except Exception:
            continue
    if len(account_header) == 1:
        return account_header[0]
    if len(account_header) > 1:
        raise RuntimeError('检测到多个账号头像，无法确认当前登录账号，请仅保留目标 QQ 主窗口。')
    # 兼容没有“在线状态”按钮的旧版：只接受主窗口左上方唯一的头像。
    rect = window.rectangle()
    header_candidates = []
    for control in candidates:
        area = control.rectangle()
        if (rect.left <= area.left < rect.left + min(180, (rect.right - rect.left) * 0.3)
                and rect.top <= area.top < rect.top + min(180, (rect.bottom - rect.top) * 0.25)):
            header_candidates.append(control)
    return header_candidates[0] if len(header_candidates) == 1 else None


def find_own_profile_account(window, desktop=None):
    """只读当前进程的本人资料卡；兼容独立窗口及分开的 QQ 标签和号码。"""
    if desktop is None:
        from pywinauto import Desktop
        desktop = Desktop(backend='uia')
    process_id = window.process_id()
    profile_names = {'资料卡', '个人资料', '我的资料'}
    combined_pattern = re.compile(r'QQ\s*(?:号|号码)?\s*[:：]?\s*([0-9]{5,12})', re.IGNORECASE)
    label_pattern = re.compile(r'QQ\s*(?:号|号码)?\s*[:：]?', re.IGNORECASE)
    accounts = set()
    for top in desktop.windows():
        try:
            if top.process_id() != process_id or not top.is_visible():
                continue
            scopes = [control for control in [top, *top.descendants()]
                      if control.is_visible() and control.window_text().strip() in profile_names]
        except Exception:
            continue
        for scope in scopes:
            try:
                controls = [control for control in scope.descendants() if control.is_visible()]
                if not any(control.element_info.control_type == 'Button'
                           and control.window_text().strip() in {'编辑资料', '修改资料'}
                           for control in controls):
                    continue
            except Exception:
                continue
            for control in controls:
                try:
                    if control.element_info.control_type not in {'Text', 'Edit'}:
                        continue
                    name = control.window_text().strip()
                    match = combined_pattern.fullmatch(name)
                    if match:
                        accounts.add(match.group(1))
                        continue
                    if not label_pattern.fullmatch(name):
                        continue
                    label_rect = control.rectangle()
                    parent = control.parent()
                    if parent is None:
                        continue
                    for sibling in parent.children():
                        if (not sibling.is_visible()
                                or sibling.element_info.control_type not in {'Text', 'Edit'}):
                            continue
                        number = sibling.window_text().strip()
                        if not re.fullmatch(r'[0-9]{5,12}', number):
                            continue
                        rect = sibling.rectangle()
                        overlap = min(label_rect.bottom, rect.bottom) - max(label_rect.top, rect.top)
                        height = min(label_rect.bottom - label_rect.top, rect.bottom - rect.top)
                        # QQ 当前版本的两块文字边界有 1 像素重叠。只接受同一行右侧的号码。
                        if (height > 0 and overlap >= height * 0.5
                                and -3 <= rect.left - label_rect.right <= max(80, height * 6)):
                            accounts.add(number)
                except Exception:
                    continue
    if len(accounts) > 1:
        raise RuntimeError('同一 QQ 进程检测到多个不同的本人账号，无法确认账号，请关闭多余资料卡后重试。')
    return next(iter(accounts), None)


# 复用已打开的本人资料卡，必要时只点击一次头像；此处不搜索或添加好友。
def read_qq_account_from_avatar(window, pyautogui, win32con, win32gui, desktop=None):
    if desktop is None:
        from pywinauto import Desktop
        desktop = Desktop(backend='uia')
    account = find_own_profile_account(window, desktop)
    if account is not None:
        return account
    activate_window(window, win32con, win32gui)
    avatar = find_own_avatar(window)
    if avatar is None:
        raise RuntimeError(f'QQ 进程 {window.process_id()} 中未找到本人头像，请显示主窗口左上角账号栏。')

    avatar.click_input()
    deadline = time.monotonic() + AVATAR_TIMEOUT_SECONDS
    while time.monotonic() < deadline:
        account = find_own_profile_account(window, desktop)
        if account is not None:
            return account
        time.sleep(0.2)
    raise RuntimeError(
        f'点击 QQ 进程 {window.process_id()} 的头像后，未能读取本人资料卡中的 QQ 号。'
        '请确认资料卡已打开并显示“编辑资料”和 QQ 号码。'
    )


# 将 ODS 单元格内容统一转换为字符串，整数数值不保留小数点。
def normalize_cell_value(cell, extract_text):
    value = cell.getAttribute("stringvalue")
    if value is None:
        value = cell.getAttribute("value")
    if value is None:
        value = extract_text(cell)

    value = str(value or "").strip()
    if not value:
        return ""

    if cell.getAttribute("valuetype") == "float":
        try:
            number = Decimal(value)
            if number == number.to_integral_value():
                return format(number, "f")
        except InvalidOperation:
            pass
    return value


# 根据重复列数定位指定列的单元格，供读取 B 列完成标记使用。
def cell_at_column(row, column_index, TableCell):
    column = 0
    for cell in row.getElementsByType(TableCell):
        repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        if column <= column_index < column + repeat:
            return cell
        column += repeat
    return None


def read_excel_number_rows(number_file):
    """读取采集工具使用的活动工作表，保留公式、格式和其他工作表。"""
    from openpyxl import load_workbook

    document = load_workbook(number_file)
    try:
        sheet = document.active
        if sheet is None or not hasattr(sheet, 'iter_rows'):
            raise ValueError('Excel 当前选中的不是数据工作表，请在 Excel 中选中号码表后保存。')
        entries = {}
        for row_index, (number_cell, status_cell) in enumerate(
            sheet.iter_rows(min_col=1, max_col=2), start=1
        ):
            value = number_cell.value
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            if number_cell.data_type == 'f':
                raise ValueError(f'第 {row_index} 行 A 列是公式，请粘贴为实际 QQ 号码值。')
            if isinstance(value, float) and value.is_integer():
                number = str(int(value))
            else:
                number = str(value).strip()
            if not re.fullmatch(r'[0-9]{5,12}', number):
                raise ValueError(f'第 {row_index} 行 A 列不是有效 QQ 号，请删除表头或修正号码。')
            entry = entries.setdefault(number, {'rows': [], 'pending': False})
            entry['rows'].append(row_index)
            if str(status_cell.value or '').strip() != PROCESSED_MARKER:
                entry['pending'] = True
        numbers = [(number, entry['rows']) for number, entry in entries.items() if entry['pending']]
        return numbers, document, sheet, number_file
    except BaseException:
        document.close()
        raise


# 读取 A 列号码及同一行 B 列状态，返回尚未标记的号码和其原始行号。
# 同一号码若重复出现，保留原有去重行为，并记录所有对应行号。
def read_number_rows(load_ods, Table, TableRow, TableCell, extract_text):
    number_file = Path(NUMBER_FILE_PATH)
    if not number_file.is_absolute():
        number_file = WORKSPACE_DIR / number_file
    if not number_file.is_file():
        raise FileNotFoundError(f"找不到号码文件：{number_file}")
    if number_file.suffix.lower() == '.xlsx':
        return read_excel_number_rows(number_file)
    if number_file.suffix.lower() != '.ods':
        raise ValueError('请选择 Excel（.xlsx）或 ODS（.ods）号码表。')

    document = load_ods(str(number_file))
    sheets = document.spreadsheet.getElementsByType(Table)
    if not sheets:
        raise RuntimeError(f"ODS 文件中没有工作表：{number_file}")

    sheet = sheets[0]
    entries = {}
    row_index = 0
    for row in sheet.getElementsByType(TableRow):
        first_index = row_index + 1
        repeat = int(row.getAttribute("numberrowsrepeated") or 1)
        row_index += repeat
        cells = row.getElementsByType(TableCell)
        if not cells:
            continue
        number = normalize_cell_value(cells[0], extract_text)
        if not number:
            continue
        if not re.fullmatch(r"[0-9]{5,12}", number):
            raise ValueError(f"第 {first_index} 行 A 列不是有效 QQ 号，请删除表头或修正号码。")
        status_cell = cell_at_column(row, 1, TableCell)
        is_processed = (
            status_cell is not None
            and normalize_cell_value(status_cell, extract_text) == PROCESSED_MARKER
        )
        if number not in entries:
            entries[number] = {"rows": [], "pending": False}
        entries[number]["rows"].extend(range(first_index, row_index + 1))
        if not is_processed:
            entries[number]["pending"] = True

    numbers = [
        (number, entry["rows"])
        for number, entry in entries.items()
        if entry["pending"]
    ]
    return numbers, document, sheet, number_file


# 在 ODS 节点中紧跟指定元素插入拆分出的重复单元格或行。
def insert_after(parent, existing, new_element):
    children = parent.childNodes
    index = children.index(existing)
    if index + 1 < len(children):
        parent.insertBefore(new_element, children[index + 1])
    else:
        parent.appendChild(new_element)


# ODS 会压缩连续空行/空列；写入单个位置前先拆分重复元素以保留其他区域。
def split_repeated_element(parent, element, first_count, repeat_attribute):
    repeat = int(element.getAttribute(repeat_attribute) or 1)
    if first_count >= repeat:
        return None
    remainder = deepcopy(element)
    remaining_count = repeat - first_count
    if first_count == 1:
        element.removeAttribute(repeat_attribute)
    else:
        element.setAttribute(repeat_attribute, str(first_count))
    if remaining_count == 1:
        remainder.removeAttribute(repeat_attribute)
    else:
        remainder.setAttribute(repeat_attribute, str(remaining_count))
    insert_after(parent, element, remainder)
    return remainder


# 将压缩存储的重复行展开到指定逻辑行，并返回目标行节点。
def row_at_index(sheet, row_index, TableRow):
    current_index = 1
    for row in list(sheet.getElementsByType(TableRow)):
        repeat = int(row.getAttribute("numberrowsrepeated") or 1)
        offset = row_index - current_index
        if 0 <= offset < repeat:
            target_row = row
            if offset:
                target_row = split_repeated_element(
                    sheet, row, offset, "numberrowsrepeated"
                )
            remaining = int(target_row.getAttribute("numberrowsrepeated") or 1)
            if remaining > 1:
                split_repeated_element(
                    sheet, target_row, 1, "numberrowsrepeated"
                )
            return target_row
        current_index += repeat
    raise IndexError(f"ODS 工作表中找不到第 {row_index} 行。")


# 将压缩存储的重复列拆分，并返回可以单独写入的指定列单元格。
def writable_cell_at_column(row, column_index, TableCell):
    current_column = 0
    for cell in list(row.getElementsByType(TableCell)):
        repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        offset = column_index - current_column
        if 0 <= offset < repeat:
            target_cell = cell
            if offset:
                target_cell = split_repeated_element(
                    row, cell, offset, "numbercolumnsrepeated"
                )
            remaining = int(target_cell.getAttribute("numbercolumnsrepeated") or 1)
            if remaining > 1:
                split_repeated_element(
                    row, target_cell, 1, "numbercolumnsrepeated"
                )
            return target_cell
        current_column += repeat

    while current_column <= column_index:
        cell = TableCell()
        row.appendChild(cell)
        current_column += 1
    return cell


# 在号码所在行的 B 列写入“已处理”，并立即保存到号码文件。
def mark_rows_processed(
    document, sheet, row_indexes, TableRow, TableCell, Paragraph, number_file
):
    if Path(number_file).suffix.lower() == '.xlsx':
        for row_index in row_indexes:
            sheet.cell(row=row_index, column=2).value = PROCESSED_MARKER
        try:
            atomic_save(document, number_file)
        except PermissionError as error:
            raise PermissionError('无法保存 Excel 号码表，请关闭 Excel/WPS 中的该文件，并确认文件夹可写；最后一条可能尚未记录，请检查日志。') from error
        return
    for row_index in row_indexes:
        row = row_at_index(sheet, row_index, TableRow)
        cell = writable_cell_at_column(row, 1, TableCell)
        cell.setAttribute("valuetype", "string")
        cell.setAttribute("stringvalue", PROCESSED_MARKER)
        if cell.getAttribute("value") is not None:
            cell.removeAttribute("value")
        for child in list(cell.childNodes):
            cell.removeChild(child)
        cell.addElement(Paragraph(text=PROCESSED_MARKER))
    atomic_save(document, number_file, ods=True)


# 获取搜索框和“搜索”按钮；必要控件缺失时停止以避免错误操作。
def get_search_controls(page):
    controls = page.descendants()
    inputs = [control for control in controls
              if control.element_info.control_type == 'Edit'
              and control.is_visible() and control.is_enabled()]
    named_inputs = [control for control in inputs if control.window_text().strip() == '搜索']
    input_candidates = named_inputs or inputs
    search_input = input_candidates[0] if len(input_candidates) == 1 else None
    search_button = next(
        (
            control
            for control in controls
            if control.element_info.control_type == "Button"
            and control.window_text().strip() == "搜索"
            and control.is_visible()
            and control.is_enabled()
        ),
        None,
    )
    if search_input is None or search_button is None:
        raise RuntimeError('QQ 搜索页面中未找到唯一可用的搜索框或搜索按钮，请保持“综合搜索/全网搜索”页面显示。')
    return search_input, search_button


# 尝试切换到“用户”搜索分类，返回是否成功。
def select_user_tab(page):
    tab = next(
        (
            control
            for control in page.descendants()
            if control.element_info.control_type == "TabItem"
            and control.window_text().strip() == USER_TAB_TEXT
            and control.is_visible()
            and control.is_enabled()
        ),
        None,
    )
    if tab is None:
        return False
    try:
        tab.click_input()
        time.sleep(0.3)
        return True
    except Exception:
        return False


# 汇总搜索页面可访问文本，用于判断搜索前后的内容变化及页面状态。
def read_search_page_text(page):
    # QQ 的 Document 名称可能为空或固定，结果文字实际位于其后代控件。
    texts = []
    seen = set()
    for control in page.descendants():
        try:
            if not control.is_visible():
                continue
            text = control.window_text().strip()
            if text and text not in seen:
                seen.add(text)
                texts.append(text)
        except Exception:
            continue
    return '\n'.join(texts)


# 判断搜索区域下方是否出现结果按钮，不限定按钮文字或可用状态。
def has_result_button(page):
    controls = page.descendants()
    search_button = next(
        (
            control
            for control in controls
            if control.element_info.control_type == "Button"
            and control.window_text().strip() == "搜索"
            and control.is_visible()
        ),
        None,
    )
    if search_button is None:
        return False

    result_top = search_button.rectangle().bottom
    return any(
        control.element_info.control_type == "Button"
        and control.window_text().strip()
        and control.window_text().strip() != "搜索"
        and control.is_visible()
        and control.rectangle().top >= result_top
        for control in controls
    )


def invoke_page_button(page, button):
    """Use the QQ button's UIA action; a covered screen coordinate is unreliable."""
    from pywinauto.uia_defines import NoPatternInterfaceError
    try:
        action = button.iface_invoke
    except (AttributeError, NoPatternInterfaceError):
        page.set_focus()
        button.click_input()
    else:
        # Never retry with a mouse click after Invoke: it may already have acted.
        action.Invoke()


def result_matches_number(page, number):
    """Match the number in result text below Search, never in the query input."""
    search_buttons = [control for control in page.descendants()
                      if control.element_info.control_type == 'Button'
                      and control.window_text().strip() == '搜索'
                      and control.is_visible() and control.is_enabled()]
    if len(search_buttons) != 1:
        return False
    search_button = search_buttons[0]
    result_top = search_button.rectangle().bottom
    pattern = re.compile(r'(?<![0-9])' + re.escape(number) + r'(?![0-9])')
    for control in page.descendants():
        try:
            if (control.element_info.control_type == 'Text'
                    and control.is_visible() and control.rectangle().top >= result_top
                    and pattern.search(control.window_text())):
                return True
        except Exception:
            continue
    return False


# 核对本次号码后，只触发结果区域内唯一的添加按钮。
def click_result_add_button(page, expected_number=None):
    page.set_focus()
    if expected_number is not None and not result_matches_number(page, expected_number):
        raise RuntimeError(f'当前搜索结果尚未显示 QQ 号 {expected_number}，未点击添加。')
    _, search_button = get_search_controls(page)
    result_top = search_button.rectangle().bottom
    buttons = [control for control in page.descendants()
               if control.element_info.control_type == 'Button'
               and control.window_text().strip() in RESULT_ACTIONS
               and control.is_visible() and control.is_enabled()
               and control.rectangle().top >= result_top]
    if not buttons:
        return False
    if len(buttons) != 1:
        raise RuntimeError('搜索结果中有多个添加按钮，无法确定目标，已停止。')
    invoke_page_button(page, buttons[0])
    return True


# 仅发送刚刚填写并核对过的申请表单，避免点击聊天窗口中的发送。
def click_send_button(desktop, process_id, timeout=5, *, form=None, expected_number=None):
    if form is None:
        form = prepare_verification_form(desktop, process_id, VERIFICATION_MESSAGE,
                                         PAGE_TIMEOUT_SECONDS, expected_number=expected_number)
    return click_verified_send(form, process_id, VERIFICATION_MESSAGE, SEND_BUTTON_TEXT,
                               expected_number=expected_number, timeout=timeout)


def current_sheet_numbers(sheet, number_file, TableRow, TableCell, extract_text):
    """用于恢复记录核对；包含当前工作表中已经标记的号码。"""
    if Path(number_file).suffix.lower() == '.xlsx':
        values = (row[0].value for row in sheet.iter_rows(min_col=1, max_col=1))
        return {str(int(value)) if isinstance(value, float) and value.is_integer()
                else str(value).strip() for value in values if value is not None}
    return {normalize_cell_value(cells[0], extract_text)
            for row in sheet.getElementsByType(TableRow)
            if (cells := row.getElementsByType(TableCell))}


# 判断页面是否有明确标注为用户的搜索结果。
def has_user_result(page):
    return any(
        control.element_info.control_type == "Text"
        and control.window_text().strip() == USER_TAB_TEXT
        and control.is_visible()
        for control in page.descendants()
    )


# 判断页面是否包含除固定导航文字外的结果文本。
def has_result_content(page):
    static_texts = set(STATIC_PAGE_TEXTS) | search_page_titles() | {'频道', '表情'}
    return any(
        control.element_info.control_type == "Text"
        and control.window_text().strip()
        and control.window_text().strip() not in static_texts
        and control.is_visible()
        for control in page.descendants()
    )


# 读取当前搜索框内容，用于判断是否已显示本次号码的结果。
def current_query(search_input):
    try:
        return str(search_input.get_value()).strip()
    except Exception:
        return search_input.window_text().strip()


# 等待页面内容稳定并判断搜索结果、无结果或超时状态。
def wait_for_search_state(page, baseline, timeout, user_tab_selected, expected_number=None):
    deadline = time.monotonic() + timeout
    changed = False
    stable_text = None
    stable_since = None

    while time.monotonic() < deadline:
        text = read_search_page_text(page)
        if text != baseline:
            changed = True
        loading = any(marker in text for marker in LOADING_MARKERS)
        result_button = has_result_button(page)
        result_content = has_result_content(page)
        user_result = has_user_result(page)
        no_results = any(marker in text for marker in NO_RESULT_MARKERS)
        target_matches = expected_number is None or result_matches_number(page, expected_number)
        result_detected = result_button or (
            user_tab_selected and result_content
        ) or (
            not user_tab_selected and (user_result or result_content)
        )
        if changed and ((result_detected and target_matches) or no_results) and not loading:
            if text == stable_text:
                if time.monotonic() - stable_since >= RESULT_STABLE_SECONDS:
                    if no_results:
                        return "no_results"
                    if not user_tab_selected and not user_result:
                        return "ignored"
                    return "results"
            else:
                stable_text = text
                stable_since = time.monotonic()
        else:
            stable_text = None
            stable_since = None
        time.sleep(0.25)

    return "timeout"


# 搜索一个号码；对超时按配置重试，并返回状态及尝试次数。
def search_number(page, number, user_tab_selected):
    page.set_focus()
    search_input, search_button = get_search_controls(page)
    existing_text = read_search_page_text(page)
    if current_query(search_input) == number and result_matches_number(page, number):
        if any(marker in existing_text for marker in NO_RESULT_MARKERS):
            return "no_results", 0
        if has_result_button(page) or (
            user_tab_selected and has_result_content(page)
        ) or (
            not user_tab_selected and (has_user_result(page) or has_result_content(page))
        ):
            if not user_tab_selected and not has_user_result(page):
                return "ignored", 0
            return "results", 0

    for attempt in range(1, MAX_SEARCH_ATTEMPTS + 1):
        page.set_focus()
        search_input.set_edit_text(number)
        time.sleep(0.2)
        if current_query(search_input) != number:
            raise RuntimeError(f'搜索框中的号码与 {number} 不一致，已停止。')
        baseline = read_search_page_text(page)
        invoke_page_button(page, search_button)

        state = wait_for_search_state(
            page, baseline, SEARCH_TIMEOUT_SECONDS, user_tab_selected, expected_number=number
        )
        if state != "timeout":
            return state, attempt
        if attempt < MAX_SEARCH_ATTEMPTS:
            print(f"第 {attempt} 次等待超时，准备重试当前号码。")

    return "timeout", MAX_SEARCH_ATTEMPTS


# 主流程：识别并过滤账号，加载待处理号码，按账号分批搜索并写入完成标记。
def main():
    if sys.platform != "win32":
        raise RuntimeError("此脚本仅支持 Windows 电脑版 QQ。")
    if not isinstance(NUMBERS_PER_ACCOUNT, int) or NUMBERS_PER_ACCOUNT <= 0:
        raise ValueError("NUMBERS_PER_ACCOUNT 必须是大于 0 的整数。")
    if not isinstance(VERIFICATION_MESSAGE, str) or not VERIFICATION_MESSAGE.strip():
        raise ValueError('加好友验证信息不能为空，请在页面填写。')
    source_file = Path(NUMBER_FILE_PATH)
    if not source_file.is_absolute():
        source_file = WORKSPACE_DIR / source_file
    # 必须在读取资料卡或操作任何 QQ 窗口前检查，避免发送后才发现无法保存。
    ensure_number_file_writable(source_file)

    (
        Desktop,
        pyautogui,
        win32con,
        win32gui,
        load_ods,
        Table,
        TableRow,
        TableCell,
        extract_text,
        Paragraph,
    ) = load_dependencies()
    pyautogui.FAILSAFE = True
    desktop = Desktop(backend="uia")
    qq_windows = find_qq_windows(desktop, win32gui)
    print(f"当前检测到 {len(qq_windows)} 个已登录 QQ 主窗口。")
    if not qq_windows:
        if TARGET_PROCESS_ID:
            raise RuntimeError(
                f'没有找到已选择的 QQ 窗口（进程 {TARGET_PROCESS_ID}）。'
                '请确认 QQ 已登录且窗口已显示，再点击“选择 QQ 窗口”重新选择。'
            )
        raise RuntimeError('没有检测到已登录的 QQ 主窗口。请先登录 QQ 并显示主窗口，再开始。')

    direct_window_mode = bool(TARGET_PROCESS_ID) and len(qq_windows) == 1 and not EXCLUDED_QQ_ACCOUNTS
    accounts = []
    seen_accounts = set()
    for window in qq_windows:
        if direct_window_mode:
            accounts.append((None, window))
            print(f'使用已选择的 QQ 窗口（进程 {window.process_id()}）；未设置排除账号，无需读取资料卡。')
            continue
        account = read_qq_account_from_avatar(
            window, pyautogui, win32con, win32gui, desktop
        )
        if account in seen_accounts:
            print(f"QQ 号 {account} 对应窗口重复，已跳过。")
            continue
        seen_accounts.add(account)
        if account in EXCLUDED_QQ_ACCOUNTS:
            print(f"QQ 号 {account} 在排除名单中，本次不参与操作。")
            continue
        accounts.append((account, window))
        print(f"已确认 QQ 号：{account}（进程 {window.process_id()}）。")

    if not accounts:
        print("所有已登录 QQ 账号都在排除名单中，没有可处理账号，脚本结束。")
        return
    account_windows = [window for _, window in accounts]

    numbers, number_document, number_sheet, number_file = read_number_rows(
        load_ods, Table, TableRow, TableCell, extract_text
    )
    sheet_name = (number_sheet.title if Path(number_file).suffix.lower() == '.xlsx'
                  else number_sheet.getAttribute('name'))
    journal = AttemptJournal(number_file, sheet_name, PROCESSED_MARKER)

    def save_marker(rows):
        mark_rows_processed(number_document, number_sheet, rows,
                            TableRow, TableCell, Paragraph, number_file)

    pending_before_recovery = len(numbers)
    numbers = journal.recover(
        numbers, save_marker,
        all_numbers=current_sheet_numbers(number_sheet, number_file, TableRow, TableCell, extract_text),
    )
    if len(numbers) < pending_before_recovery:
        print(f'已补写上次未保存的 {pending_before_recovery - len(numbers)} 个号码标记，本次跳过这些号码。')
    if not numbers:
        print("没有待处理号码：号码列表为空或都已在 B 列标记为“已处理”。")
        return
    if len(numbers) < NUMBERS_PER_ACCOUNT:
        print(
            f"剩余号码不足 {NUMBERS_PER_ACCOUNT} 个，按要求不处理并结束。"
        )
        return

    no_result_count = 0
    ignored_result_count = 0
    timeout_numbers = []
    processed_count = 0
    print(f"共读取到 {len(numbers)} 个不重复号码，每个账号处理 {NUMBERS_PER_ACCOUNT} 个。")
    for account_index, (account, account_window) in enumerate(accounts, start=1):
        remaining = len(numbers) - processed_count
        if remaining < NUMBERS_PER_ACCOUNT:
            print(
                f"剩余号码 {remaining} 个，不足 {NUMBERS_PER_ACCOUNT} 个，按要求立即结束。"
            )
            break

        activate_qq_account_window(
            account_window, account_windows, win32con, win32gui
        )
        # 单窗口已在上方确认过；没有账号切换时不关闭资料卡后再次打开。
        # 多窗口仍在切换后复核，先确认再清理当前窗口的附属页面。
        if not direct_window_mode and len(qq_windows) > 1:
            confirmed_account = read_qq_account_from_avatar(
                account_window, pyautogui, win32con, win32gui, desktop
            )
            if confirmed_account != account:
                raise RuntimeError(
                    f'窗口账号确认不一致：预期 {account}，当前为 {confirmed_account}；'
                    '已停止，请重新选择正确的 QQ 窗口。'
                )
        if CLOSE_SECONDARY_WINDOWS:
            close_qq_secondary_windows(
                desktop, [account_window], win32con, win32gui
            )
        account_label = account if account is not None else f'所选窗口（进程 {account_window.process_id()}）'
        print(
            f"使用第 {account_index}/{len(accounts)} 个账号窗口 {account_label}，"
            f"处理接下来的 {NUMBERS_PER_ACCOUNT} 个号码。"
        )

        window, page = open_add_friend_page(
            desktop, win32con, win32gui, account_window
        )
        user_tab_selected = select_user_tab(page)
        if user_tab_selected:
            print("已选择 QQ 搜索页的“用户”选项。")
        else:
            print("无法识别或切换“用户”选项，将只接受明确标注为用户的结果。")

        batch_end = processed_count + NUMBERS_PER_ACCOUNT
        while processed_count < batch_end:
            number, source_rows = numbers[processed_count]
            index = processed_count + 1
            print(f"账号 {account_label} 正在搜索第 {index}/{len(numbers)} 个号码：{number}。")
            state, attempts = search_number(page, number, user_tab_selected)
            processed_count += 1
            if state == "no_results":
                no_result_count += 1
                print(f"第 {index} 个号码无搜索结果，继续下一个。")
            elif state == "ignored":
                ignored_result_count += 1
                print(f"第 {index} 个号码的结果不是明确的用户结果，已忽略并继续。")
            elif state == "timeout":
                timeout_numbers.append(number)
                print(f"第 {index} 个号码搜索超时，已重试 {attempts} 次并记录，继续下一个。")
            else:
                print(f"第 {index} 个号码的搜索结果已加载。")
                if click_result_add_button(page, expected_number=number):
                    print(f"已点击第 {index} 个号码结果中的“添加”按钮。")
                    try:
                        form = prepare_verification_form(
                            desktop, window.process_id(), VERIFICATION_MESSAGE, PAGE_TIMEOUT_SECONDS,
                            expected_number=number,
                        )
                    except NoVerificationField as missing_field:
                        cancel_unconfigured_form(
                            missing_field.window, window.process_id(), number, SEND_TIMEOUT_SECONDS
                        )
                        print(
                            f'QQ 号 {number} 的申请页没有验证信息输入框；已取消，未发送且未标记。'
                            '继续处理下一个号码。'
                        )
                        time.sleep(NUMBER_INTERVAL_SECONDS)
                        continue
                    print('已填写并核对页面配置的验证信息。')
                    if not AUTO_SEND:
                        print("自动发送已关闭，已填写申请表单；本次在此停止，该号码未标记。")
                        return
                    # 发送前落盘；即使此刻按停止，也不能在下次启动时无记录重发。
                    journal.record(number, source_rows, 'sending')
                    if not click_send_button(desktop, window.process_id(), SEND_TIMEOUT_SECONDS,
                                             form=form, expected_number=number):
                        raise RuntimeError('未能点击申请表单中的发送，请核对该号码的申请状态。')
                    print(f'QQ 号 {number}：已触发“发送”并确认申请窗口关闭，不代表对方已通过。')
                else:
                    print(f"第 {index} 个号码结果已处理；没有可点击的“添加”按钮，跳过。")

            journal.record(number, source_rows, 'processed')
            try:
                save_marker(source_rows)
            except Exception as error:
                raise RuntimeError(
                    f'号码表保存失败：{error} 已保留 QQ 号 {number} 的本地处理记录；'
                    f'关闭 Excel/WPS 中的号码表后重新开始，将先补写 B 列，再继续后续号码。'
                    f'恢复记录：{journal.path}'
                ) from error
            journal.clear(number)
            print(f"已在号码所在行的 B 列标记“{PROCESSED_MARKER}”。")
            time.sleep(NUMBER_INTERVAL_SECONDS)

    print(
        f"搜索完成：处理 {processed_count}/{len(numbers)} 个，"
        f"明确无结果 {no_result_count} 个，"
        f"忽略非用户或类型不明结果 {ignored_result_count} 个，"
        f"超时 {len(timeout_numbers)} 个。"
    )
    if timeout_numbers:
        print("超时号码：" + ", ".join(timeout_numbers))
    print("脚本结束；QQ 保持在最后一个号码的搜索页面。")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"错误：{error}", file=sys.stderr)
        raise SystemExit(1) from error
