"""Append captured QQ numbers to the ODS input without resetting existing status."""

from datetime import datetime
from pathlib import Path
import os
import re
import shutil
import tempfile

from odf.opendocument import OpenDocumentSpreadsheet, load
from odf.table import Table, TableCell, TableRow
from odf.teletype import extractText
from odf.text import P
from openpyxl import load_workbook

from open_qq_add_friend import NUMBER_FILE_PATH, WORKSPACE_DIR
from zq_qq import NUMBER_WORKBOOK_PATH


def sync_numbers(source=NUMBER_WORKBOOK_PATH, destination=NUMBER_FILE_PATH):
    source = Path(source)
    destination = Path(destination)
    if not destination.is_absolute():
        destination = WORKSPACE_DIR / destination
    numbers = []
    invalid = 0
    if source.is_file():
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            for (value,) in workbook.active.iter_rows(min_col=1, max_col=1, values_only=True):
                if value is None:
                    continue
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                text = str(value).strip()
                if re.fullmatch(r"[0-9]{5,12}", text):
                    numbers.append(text)
                elif text:
                    invalid += 1
        finally:
            workbook.close()
    existed = destination.is_file()
    if existed:
        document = load(str(destination))
        sheets = document.spreadsheet.getElementsByType(Table)
        if not sheets:
            raise ValueError("号码 ODS 中没有工作表，请检查文件。")
        sheet = sheets[0]
    else:
        document = OpenDocumentSpreadsheet()
        sheet = Table(name="QQ号码")
        document.spreadsheet.addElement(sheet)
    existing = set()
    for row in sheet.getElementsByType(TableRow):
        cells = row.getElementsByType(TableCell)
        if not cells:
            continue
        cell = cells[0]
        value = cell.getAttribute("stringvalue")
        if value is None:
            value = cell.getAttribute("value")
        if value is None:
            value = extractText(cell)
        text = str(value or "").strip()
        if re.fullmatch(r"[0-9]+[.]0+", text):
            text = text.split(".", 1)[0]
        if text:
            existing.add(text)
    new_numbers = list(dict.fromkeys(number for number in numbers if number not in existing))
    for number in new_numbers:
        row = TableRow()
        cell = TableCell(valuetype="string", stringvalue=number)
        cell.addElement(P(text=number))
        row.addElement(cell)
        row.addElement(TableCell())
        sheet.addElement(row)
    backup = None
    if new_numbers or not existed:
        destination.parent.mkdir(parents=True, exist_ok=True)
        if existed:
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            backup = destination.with_name(f"{destination.stem}.{stamp}.bak.ods")
            shutil.copy2(destination, backup)
        descriptor, temp_name = tempfile.mkstemp(prefix=".qq_numbers_", suffix=".ods", dir=destination.parent)
        os.close(descriptor)
        try:
            document.save(temp_name, addsuffix=False)
            os.replace(temp_name, destination)
        finally:
            if os.path.exists(temp_name):
                os.unlink(temp_name)
    return {"source": str(source), "source_exists": source.is_file(),
            "destination": str(destination), "added": len(new_numbers),
            "invalid_skipped": invalid, "backup": str(backup) if backup else None}


if __name__ == "__main__":
    try:
        result = sync_numbers()
        print(f"号码文件：{result['destination']}")
        print(f"本次新增 {result['added']} 个号码，已有号码及 B 列状态保持原样。")
        if not result['source_exists']:
            print("尚未发现抓取结果。已准备好空的号码表；抓取完成后再运行此入口。")
        if result['invalid_skipped']:
            print(f"跳过 {result['invalid_skipped']} 个非 5～12 位纯数字内容，请检查源表。")
        if result['backup']:
            print(f"原号码表已备份：{result['backup']}")
    except Exception as error:
        print(f"转换失败：{error}")
        raise SystemExit(1) from error
