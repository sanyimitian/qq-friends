"""Offline regression coverage for sharing capture XLSX files with the add app."""
from pathlib import Path
import os
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from odf.opendocument import load as load_ods
from odf.table import Table, TableCell, TableRow
from odf.teletype import extractText
from odf.text import P
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

import gui_config
import desktop_app
import open_qq_add_friend as adding
import zq_qq as capture


class ExcelInputTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='qq_excel_input_test_')
        self.directory = Path(self.temporary.name)
        self.path = self.directory / 'numbers.xlsx'

    def tearDown(self):
        self.temporary.cleanup()

    def write_rows(self, rows, path=None):
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        workbook.save(path or self.path)
        workbook.close()

    def read_numbers(self, path=None):
        with patch.object(adding, 'NUMBER_FILE_PATH', path or self.path):
            result = adding.read_number_rows(
                load_ods, Table, TableRow, TableCell, extractText
            )
        self.addCleanup(result[1].close)
        return result

    def mark(self, document, sheet, rows, path=None):
        adding.mark_rows_processed(
            document, sheet, rows, TableRow, TableCell, P, path or self.path
        )

    def test_gui_default_matches_capture_excel_default(self):
        add_default = gui_config.defaults('add')['NUMBER_FILE_PATH']
        capture_default = gui_config.defaults('capture')['NUMBER_WORKBOOK_PATH']
        self.assertEqual(Path(add_default), Path(capture_default))
        self.assertEqual(Path(add_default).suffix, '.xlsx')

    def test_gui_accepts_excel_and_legacy_ods_and_sets_path_type(self):
        for suffix in ('.xlsx', '.XLSX', '.ods', '.ODS'):
            with self.subTest(suffix=suffix):
                path = self.directory / ('numbers' + suffix)
                config = gui_config.defaults('add')
                config['NUMBER_FILE_PATH'] = str(path)
                validated = gui_config.validate('add', config)
                self.assertEqual(Path(validated['NUMBER_FILE_PATH']), path.resolve())
                module = SimpleNamespace()
                gui_config.apply_config('add', module, validated)
                self.assertIsInstance(module.NUMBER_FILE_PATH, Path)
                self.assertEqual(module.NUMBER_FILE_PATH, path.resolve())

    def test_gui_rejects_unsupported_extension_and_missing_input(self):
        for name in ('numbers.xls', 'numbers.csv', 'numbers.txt'):
            with self.subTest(name=name):
                config = gui_config.defaults('add')
                config['NUMBER_FILE_PATH'] = str(self.directory / name)
                with self.assertRaises(ValueError):
                    gui_config.validate('add', config)
        config = gui_config.defaults('add')
        config['NUMBER_FILE_PATH'] = str(self.path)
        with self.assertRaises(ValueError):
            gui_config.validate('add', config, for_run=True)

    def test_excel_chooser_binds_and_saves_source_path_without_conversion(self):
        self.write_rows([['123456789']])
        original = self.path.read_bytes()
        value = {'path': str(self.directory / 'old.ods')}
        variable = SimpleNamespace(
            get=lambda: value['path'],
            set=lambda new: value.update(path=new),
        )
        app = SimpleNamespace(
            root=None, mode='add', inputs={'NUMBER_FILE_PATH': variable},
            collect=lambda: {'NUMBER_FILE_PATH': value['path']},
            append_log=Mock(), status=Mock(),
        )
        with patch('tkinter.filedialog.askopenfilename', return_value=str(self.path)), \
             patch('tkinter.messagebox.showerror') as error, \
             patch('desktop_app.save_config') as save:
            desktop_app.App.choose_excel(app)
        self.assertEqual(value['path'], str(self.path))
        save.assert_called_once_with('add', {'NUMBER_FILE_PATH': str(self.path)})
        error.assert_not_called()
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.directory.iterdir()), [self.path])

    def test_cancelling_excel_chooser_preserves_selected_path(self):
        variable = Mock()
        variable.get.return_value = str(self.directory / 'old.ods')
        app = SimpleNamespace(root=None, inputs={'NUMBER_FILE_PATH': variable}, collect=Mock())
        with patch('tkinter.filedialog.askopenfilename', return_value=''):
            desktop_app.App.choose_excel(app)
        variable.set.assert_not_called()
        app.collect.assert_not_called()

    def test_reads_integer_text_blank_and_duplicate_rows_in_source_order(self):
        self.write_rows([
            [123456789],
            [' 234567890 '],
            [],
            ['123456789', adding.PROCESSED_MARKER],
            ['345678901', adding.PROCESSED_MARKER],
            ['456789012', adding.PROCESSED_MARKER],
            ['456789012'],
            ['00012345'],
            [567890123.0],
        ])
        numbers, _, _, path = self.read_numbers()
        self.assertEqual(numbers, [
            ('123456789', [1, 4]),
            ('234567890', [2]),
            ('456789012', [6, 7]),
            ('00012345', [8]),
            ('567890123', [9]),
        ])
        self.assertEqual(path, self.path)

    def test_invalid_qq_values_fail_with_the_source_row(self):
        invalid_values = (
            'QQ number', '1234', '1234567890123', '12345.5',
            12345.5, True, '=123456789', '１２３４５６', '123 456',
        )
        for value in invalid_values:
            with self.subTest(value=value):
                self.write_rows([['123456789'], [value]])
                with self.assertRaises(ValueError) as raised:
                    self.read_numbers()
                self.assertIn('2', str(raised.exception))

    def test_duplicate_rows_are_all_marked_and_skipped_after_reopen(self):
        self.write_rows([
            ['123456789'],
            ['234567890', adding.PROCESSED_MARKER],
            ['123456789', adding.PROCESSED_MARKER],
        ])
        numbers, document, sheet, _ = self.read_numbers()
        self.assertEqual(numbers, [('123456789', [1, 3])])
        self.mark(document, sheet, numbers[0][1])
        self.assertEqual(self.read_numbers()[0], [])
        reloaded = load_workbook(self.path)
        try:
            self.assertEqual(
                [reloaded.active.cell(row, 2).value for row in (1, 2, 3)],
                [adding.PROCESSED_MARKER] * 3,
            )
        finally:
            reloaded.close()
        self.assertEqual(list(self.directory.glob('.qq_save_*')), [])

    def test_active_sheet_and_unrelated_content_survive_marker_write(self):
        workbook = Workbook()
        cover = workbook.active
        cover.title = 'Notes'
        cover['A1'] = 'Keep this cover sheet'
        sheet = workbook.create_sheet('Numbers')
        workbook.active = 1
        sheet['A1'] = '123456789'
        sheet['A2'] = '234567890'
        sheet['B1'].font = Font(name='Calibri', bold=True, color='0070C0')
        sheet['B1'].fill = PatternFill('solid', fgColor='FFFF00')
        sheet['C1'] = '=SUM(2,3)'
        sheet['D1'] = 'Retain my note'
        sheet['A1'].number_format = '@'
        sheet.column_dimensions['C'].width = 24
        sheet.freeze_panes = 'A2'
        workbook.save(self.path)
        workbook.close()
        numbers, document, active, _ = self.read_numbers()
        self.assertEqual(active.title, 'Numbers')
        self.assertEqual(numbers, [('123456789', [1]), ('234567890', [2])])
        self.mark(document, active, [1])
        check = load_workbook(self.path)
        try:
            self.assertEqual(check.sheetnames, ['Notes', 'Numbers'])
            self.assertEqual(check.active.title, 'Numbers')
            self.assertEqual(check['Notes']['A1'].value, 'Keep this cover sheet')
            self.assertIsNone(check['Notes']['B1'].value)
            self.assertEqual(check.active['A1'].value, '123456789')
            self.assertEqual(check.active['A1'].number_format, '@')
            self.assertEqual(check.active['B1'].value, adding.PROCESSED_MARKER)
            self.assertTrue(check.active['B1'].font.bold)
            self.assertEqual(check.active['B1'].fill.fgColor.rgb, '00FFFF00')
            self.assertEqual(check.active['C1'].value, '=SUM(2,3)')
            self.assertEqual(check.active['D1'].value, 'Retain my note')
            self.assertEqual(check.active.column_dimensions['C'].width, 24)
            self.assertEqual(check.active.freeze_panes, 'A2')
        finally:
            check.close()
        self.assertEqual(self.read_numbers()[0], [('234567890', [2])])

    def test_custom_processed_marker_is_used_for_reading_and_writing(self):
        self.write_rows([['123456789', 'done'], ['234567890']])
        with patch.object(adding, 'PROCESSED_MARKER', 'done'):
            numbers, document, sheet, _ = self.read_numbers()
            self.assertEqual(numbers, [('234567890', [2])])
            self.mark(document, sheet, [2])
            self.assertEqual(self.read_numbers()[0], [])

    def test_failed_atomic_replace_keeps_original_file_and_removes_temporary(self):
        self.write_rows([['123456789']])
        original = self.path.read_bytes()
        _, document, sheet, _ = self.read_numbers()
        with patch('task_control.os.replace', side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):
                self.mark(document, sheet, [1])
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.directory.glob('.qq_save_*')), [])
        self.assertEqual(self.read_numbers()[0], [('123456789', [1])])

    def test_capture_append_preserves_existing_status_and_duplicate_skip(self):
        self.write_rows([['123456789']])
        _, document, sheet, _ = self.read_numbers()
        self.mark(document, sheet, [1])
        document.close()
        with patch.object(capture, 'NUMBER_WORKBOOK_PATH', self.path):
            workbook, worksheet, existing = capture.open_number_workbook(Workbook, load_workbook)
            try:
                self.assertFalse(capture.append_number(workbook, worksheet, existing, '123456789'))
                self.assertTrue(capture.append_number(workbook, worksheet, existing, '234567890'))
            finally:
                workbook.close()
        numbers, _, active, _ = self.read_numbers()
        self.assertEqual(numbers, [('234567890', [2])])
        self.assertEqual(active['B1'].value, adding.PROCESSED_MARKER)
        self.assertIsNone(active['B2'].value)
        self.assertEqual(active.max_row, 2)

    def test_failed_save_recovers_real_excel_by_number_after_rows_move(self):
        self.write_rows([['123456789'], ['234567890']])
        with patch.dict(os.environ, {'QQ_ASSISTANT_STATE_DIR': str(self.directory / 'state')}):
            numbers, document, sheet, _ = self.read_numbers()
            journal = adding.AttemptJournal(self.path, sheet.title)
            journal.record('123456789', [1], 'sending')
            journal.record('123456789', [1], 'processed')
            with patch('task_control.os.replace', side_effect=PermissionError('locked')):
                with self.assertRaises(PermissionError):
                    self.mark(document, sheet, [1])
            self.assertEqual(self.read_numbers()[0], numbers)
            self.write_rows([['234567890'], ['123456789']])
            pending, document, sheet, _ = self.read_numbers()
            all_numbers = adding.current_sheet_numbers(sheet, self.path, TableRow, TableCell, extractText)
            recovered = adding.AttemptJournal(self.path, sheet.title).recover(
                pending, lambda rows: self.mark(document, sheet, rows), all_numbers)
            self.assertEqual(recovered, [('234567890', [1])])
            pending, _, check, _ = self.read_numbers()
            self.assertEqual(pending, [('234567890', [1])])
            self.assertIsNone(check['B1'].value)
            self.assertEqual(check['B2'].value, adding.PROCESSED_MARKER)
            self.assertFalse(journal.path.exists())


if __name__ == '__main__':
    unittest.main(verbosity=2)
